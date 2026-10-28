"""R51: el escritor propio de .xlsx produce libros válidos y rápidos."""
import math
import time

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from core.utils.xlsx_writer import escribir_xlsx, nombre_de_hoja


def test_ida_y_vuelta_con_openpyxl(tmp_path):
    ruta = tmp_path / "a.xlsx"
    escribir_xlsx(ruta, {"Datos": (["x", "y", "nota"], [[1, 2.5, "hola"], [2, -0.125, "a & b <c>"], [3, 1e-9, ""]])})
    ws = load_workbook(ruta)["Datos"]
    assert [c.value for c in ws[1]] == ["x", "y", "nota"]
    assert [c.value for c in ws[2]] == [1, 2.5, "hola"]
    assert ws.cell(3, 3).value == "a & b <c>"
    assert ws.cell(4, 2).value == 1e-9


def test_varias_hojas_nombres_validos_y_sin_repetir(tmp_path):
    ruta = tmp_path / "b.xlsx"
    escribir_xlsx(ruta, {"Parámetros": (["k"], [["v"]]), "Mapa [Hz/s]": (["a"], [[1]]), "Mapa (Hz/s)": (["a"], [[2]]),
                         "x" * 40: (["a"], [[3]])})
    nombres = load_workbook(ruta).sheetnames
    assert len(nombres) == 4 and len({n.lower() for n in nombres}) == 4
    assert all(len(n) <= 31 and not any(c in n for c in "[]:*?/\\") for n in nombres)


def test_nan_e_infinitos_quedan_vacios_y_los_tipos_de_numpy_se_escriben(tmp_path):
    ruta = tmp_path / "c.xlsx"
    escribir_xlsx(ruta, {"H": (["a", "b", "c", "d"], [[float("nan"), math.inf, np.float64(1.5), np.int64(7)],
                                                       [np.bool_(True), None, False, "x"]])})
    ws = load_workbook(ruta)["H"]
    assert ws.cell(2, 1).value is None and ws.cell(2, 2).value is None
    assert ws.cell(2, 3).value == 1.5 and ws.cell(2, 4).value == 7
    assert ws.cell(3, 1).value is True and ws.cell(3, 3).value is False and ws.cell(3, 4).value == "x"


def test_filas_de_distinto_largo_no_rompen(tmp_path):
    ruta = tmp_path / "d.xlsx"
    escribir_xlsx(ruta, {"H": (["a", "b"], [[1, 2], [3], [4, 5]])})
    df = pd.read_excel(ruta)
    assert df.shape == (3, 2) and pd.isna(df.iloc[1, 1]) and df.iloc[2, 1] == 5


def test_caracteres_de_control_no_invalidan_el_xml(tmp_path):
    ruta = tmp_path / "e.xlsx"
    escribir_xlsx(ruta, {"H": (["a"], [["mal\x00\x08texto"]])})
    assert load_workbook(ruta)["H"].cell(2, 1).value == "maltexto"


def test_nombre_de_hoja_unico():
    usados = set()
    assert nombre_de_hoja("A", usados) == "A"
    assert nombre_de_hoja("a", usados) == "a_2"


def test_cien_mil_filas_en_menos_de_dos_segundos(tmp_path):
    filas = np.column_stack([np.arange(100_000.0), np.arange(100_000.0) * 2, np.arange(100_000.0) * 3]).tolist()
    ruta = tmp_path / "f.xlsx"
    t0 = time.perf_counter()
    escribir_xlsx(ruta, {"Datos": (["x", "y", "z"], filas)})
    dt = time.perf_counter() - t0
    assert dt < 2.0, f"{dt:.2f} s"
    df = pd.read_excel(ruta)
    assert df.shape == (100_000, 3) and df.iloc[-1, 2] == 299997.0
